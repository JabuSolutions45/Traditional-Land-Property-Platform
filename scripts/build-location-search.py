"""Build Umeli's compact, read-only server search index from the master workbook.

This script does not modify the workbook. It preserves source identifiers and
source status labels so search results can explain when geography is historical
or awaiting a current source.
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = (
    ROOT
    / "outputs"
    / "location-directory-2026-08-12"
    / "Umeli_South_Africa_Location_Master_MDB_2026.xlsx"
)
OUTPUT = ROOT / "data" / "location-search.json"
SHARD_SIZE = 3_500


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def sheet_records(sheet: Any) -> Iterable[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    next(rows)
    headers = next(rows)
    for row in rows:
        if any(value is not None for value in row):
            yield dict(zip(headers, row, strict=False))


def compact_number(value: Any) -> float | int | None:
    if value is None:
        return None
    number = float(value)
    return int(number) if number.is_integer() else round(number, 7)


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    places = list(sheet_records(workbook["Places"]))
    parent_names = {
        record["place_source_id"]: record["canonical_name"] for record in places
    }

    rows: list[list[Any]] = []

    for record in places:
        rows.append(
            [
                record["place_source_id"],
                "place",
                record["canonical_name"],
                normalize(record["normalized_search_name"]),
                record["place_level"],
                record["province_name"],
                record["district_municipality_name"],
                record["municipality_name"],
                parent_names.get(record["parent_place_source_id"]),
                record["subarea_category"],
                record["primary_ward_code_2011"],
                compact_number(record["primary_ward_number_2011"]),
                record["formal_landmark_examples"],
                record["review_status"],
                compact_number(record["latitude"]),
                compact_number(record["longitude"]),
            ]
        )

    for record in sheet_records(workbook["Municipalities"]):
        rows.append(
            [
                record["municipality_source_id"],
                "municipality",
                record["municipality_name"],
                normalize(record["municipality_name"]),
                record["municipality_type"],
                record["province_name"],
                record["parent_district_name"],
                record["municipality_name"],
                None,
                record["municipality_type"],
                None,
                None,
                None,
                "historical_baseline",
                None,
                None,
            ]
        )

    for record in sheet_records(workbook["Provinces"]):
        rows.append(
            [
                record["province_source_id"],
                "province",
                record["province_name"],
                normalize(record["province_name"]),
                "province",
                record["province_name"],
                None,
                None,
                None,
                "province",
                None,
                None,
                None,
                "historical_baseline",
                None,
                None,
            ]
        )

    for record in sheet_records(workbook["Wards"]):
        ward_number = compact_number(record["ward_number"])
        municipality = record["municipality_name"]
        rows.append(
            [
                record["ward_source_id"],
                "ward",
                f"Ward {ward_number}",
                normalize(f"ward {ward_number} {municipality}"),
                "ward_2011",
                record["province_name"],
                record["district_name"],
                municipality,
                None,
                "ward_2011",
                record["ward_code"],
                ward_number,
                None,
                "historical_baseline",
                compact_number(record["latitude"]),
                compact_number(record["longitude"]),
            ]
        )

    for record in sheet_records(workbook["Traditional_Authorities"]):
        rows.append(
            [
                record["authority_source_id"],
                "traditional_authority",
                record["authority_name"],
                normalize(record["authority_name"]),
                record["authority_level"],
                record["province_name"],
                None,
                record["municipality_code"],
                record["parent_authority_source_id"],
                record["authority_level"],
                None,
                None,
                None,
                record["current_confirmation_status"],
                None,
                None,
            ]
        )

    # Named gaps are searchable so users receive an honest explanation instead
    # of a false "not found" result. No location or ward is guessed.
    for record in sheet_records(workbook["Ward_2026_Mapping_QA"]):
        if record["mapping_status"] != "pending_final_spatial_import" and record[
            "mapping_status"
        ] != "requires_current_place_source":
            continue
        rows.append(
            [
                f"qa-gap:{normalize(record['place_name']).replace(' ', '-')}",
                "data_gap",
                record["place_name"],
                normalize(record["place_name"]),
                "current_place_gap",
                record["province_name"],
                None,
                record["municipality_context"],
                None,
                "requires_current_source",
                None,
                None,
                None,
                record["mapping_status"],
                None,
                None,
            ]
        )

    shard_names: list[str] = []
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    for stale_shard in OUTPUT.parent.glob("location-search-[0-9][0-9].json"):
        stale_shard.unlink()

    for index, start in enumerate(range(0, len(rows), SHARD_SIZE)):
        shard_name = f"location-search-{index:02d}.json"
        shard_names.append(shard_name)
        (OUTPUT.parent / shard_name).write_text(
            json.dumps(
                {"rows": rows[start : start + SHARD_SIZE]},
                ensure_ascii=False,
                separators=(",", ":"),
            ),
            encoding="utf-8",
        )

    payload = {
        "version": "2026-08-12",
        "source": WORKBOOK.name,
        "columns": [
            "id",
            "kind",
            "name",
            "normalizedName",
            "level",
            "province",
            "district",
            "municipality",
            "parent",
            "category",
            "wardCode2011",
            "wardNumber2011",
            "landmarks",
            "status",
            "latitude",
            "longitude",
        ],
        "counts": {
            "places": len(places),
            "totalSearchRecords": len(rows),
        },
        "shards": shard_names,
    }

    OUTPUT.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(OUTPUT.relative_to(ROOT)),
                "places": len(places),
                "search_records": len(rows),
                "bytes": OUTPUT.stat().st_size,
                "shards": len(shard_names),
                "largest_shard_bytes": max(
                    (OUTPUT.parent / name).stat().st_size for name in shard_names
                ),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
